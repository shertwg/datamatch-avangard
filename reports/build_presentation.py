"""Сборка презентации: reports/presentation.pptx (+ reports/presentation.pdf по флагу --pdf).

Источник текста — reports/presentation_outline.md. Цифры правим в SLIDES ниже и пересобираем:

    python reports/build_presentation.py          # только .pptx
    python reports/build_presentation.py --pdf    # .pptx + .pdf (PowerPoint или LibreOffice)

Зависимость python-pptx — инструмент для отчёта, не часть решения, поэтому её нет в requirements.txt:

    pip install python-pptx

Пометки ⚠ из outline можно копировать как есть: скрипт удаляет их при сборке (и на слайдах, и в заметках).
**жирный** в тексте — жирный фрагмент. Скрипт оценивает высоту текста по метрикам Calibri и
печатает предупреждение, если текст не помещается в свой блок.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
OUT_PPTX = ROOT / "presentation.pptx"

# ---------------------------------------------------------------- стиль
FONT = "Calibri"
W, H = 13.333, 7.5                    # 16:9, дюймы
ML = MR = 0.6                         # поля слева/справа
CW = W - ML - MR                      # ширина контента
TITLE_Y, TITLE_H = 0.35, 0.85
TOP, BOTTOM = 1.35, 7.0               # область контента
GAP = 0.3
BODY_PT = 18                          # тезисы — не меньше 18 pt
TITLE_PT = 32
TABLE_PT_MIN = 14

INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0xB5, 0x12, 0x1B)   # красный «Авангарда»
TINT = RGBColor(0xF5, 0xF2, 0xF2)
LINE = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ---------------------------------------------------------------- контент
TITLE_SLIDE = {
    "title": "Прогноз спроса на домашние матчи «Авангарда»",
    "subtitle": "DataMatch, технический трек",
    "team": "Команда «shertwg»: augustteske · sharbl4 (Антон) · ArinaPig (Арина) · Алиса",
}

SLIDES = [
    {   # 1
        "layout": "text_callouts",
        "title": "Задача и метрика",
        "bullets": [
            "Прогнозируем, сколько разовых билетов продастся на каждый из 17 оставшихся домашних матчей сезона 2025/26 в каждой из 8 ценовых зон: 136 прогнозов.",
            "Прогноз — итог за всё время продаж. Дата отсечки 15.12.2025: на 6 матчей продажи уже идут, на 11 ещё не открыты.",
            "Метрика — RMSE в билетах на зону. Ориентир — базовый прогноз: среднее по зоне за прошлые матчи.",
            "Ограничения: прогноз не меньше уже проданного и не больше свободных мест (места зоны минус абонементы).",
        ],
        "callouts": [("17 × 8 = 136", "прогнозов: матч × зона"),
                     ("RMSE", "метрика, билетов на зону")],
        "notes": "«Клубу нужно знать спрос по зонам до открытия продаж, чтобы планировать цены. Поэтому мы честно имитируем момент прогноза: используем только то, что было известно на 15 декабря».",
    },
    {   # 2
        "layout": "text_top_image",
        "title": "Данные и что мы в них увидели",
        "bullets": [
            "102 домашних матча за 3 сезона: 85 сыгранных (02.09.2023–07.12.2025) для обучения и 17 для прогноза; 23 039 строк продаж по дням и зонам.",
            "Абонементы заметно сокращают свободные места: в зоне 1 из 1 141 места абонементами занято в среднем 545.",
            "Зоны заполняются очень по-разному: зоны 1, 2 и 8 — на 78–81%, а зона 6 — только на 39%, хотя свободных мест в ней столько же, сколько в зоне 7 (≈1 248), а та заполняется на 63%.",
            "Сезонность: в декабре прошлых сезонов в среднем 4 297 билетов за матч, в январе–феврале 4 889–4 936.",
        ],
        "image": "00_fill_by_zone.png",
        "notes": "«Самая интересная зона — шестая: места есть, а продаётся меньше 40%. Цен в данных нет, поэтому объяснить это мы не можем, но клубу стоит на неё посмотреть».\n\n"
                 "На графике — средняя заполняемость по зонам (среднее, не медиана).\n\n"
                 "Источники: solution.ipynb → Данные; 01_eda.ipynb (места и абонементы по зонам); 04_report.ipynb → 2.0 (заполняемость), → 1.1 (среднее по месяцам).",
    },
    {   # 3
        "layout": "text_top_image",
        "title": "Проверка без утечек",
        "bullets": [
            "Валидация повторяет реальную ситуацию: выбираем дату «псевдо-сегодня» в прошлом, учимся только на сыгранных до неё матчах и прогнозируем следующие.",
            "3 таких даты: 15.12.2023 (обучение на 20 матчах, проверка на 14), 25.11.2024 (51 → 17), 20.09.2025 (73 → 12). Всего 43 матча на проверке.",
            "Все признаки, включая продажи, результаты КХЛ и средние по соперникам, считаются только по данным до даты прогноза. Случайного перемешивания строк нет.",
            "Все модели сравниваются по одним правилам: тот же клиппинг (не меньше проданного, не больше свободных мест) и округление.",
        ],
        "image": "05_validation_scheme.png",
        "notes": "«Первая дата — тот же момент календаря, что и сейчас, вторая — ровно 17 матчей до конца сезона, как в задаче, третья — текущий сезон. Если модель хорошо работает на всех трёх, ей можно доверять на реальном прогнозе».\n\n"
                 "Источник: src/validation.py (докстринг); 04_report.ipynb → 2.5.",
    },
    {   # 4
        "layout": "text_top_diagram",
        "title": "Подход: модель и признаки",
        "bullets": [
            "Два ориентира: baseline из задания (среднее по зоне за все прошлые матчи) и «умный baseline» — среднее по зоне за текущий сезон, не ниже уже проданного.",
            "Модель — CatBoost. Каждый сыгранный матч мы превращаем в несколько «снимков» продаж: за 1, 3, …, 40 дней до игры, плюс снимок «продажи ещё не открыты». Так модель учится и на матчах в продаже, и на ещё не открытых.",
            "Признаки: уровень зоны в сезоне и за последние 5 матчей; продажи на дату (продано, за 7 и 14 дней, темп относительно типичной кривой продаж); календарь (день недели, час, месяц); соперник (история продаж с ним в Омске, его сила в КХЛ); форма «Авангарда».",
            "Главный вклад дают продажи на дату: без них RMSE вырастает до 299,3 ⚠. Вклад календаря и соперника в пределах погрешности валидации.",
        ],
        "diagram": [
            ("Снимки продаж", "за 1, 3, …, 40 дней до игры + «продажи не открыты»"),
            ("Две модели CatBoost", "матч не в продаже / матч в продаже"),
            ("Клиппинг", "[уже продано; свободные места]"),
        ],
        "notes": "«Модель предсказывает не итог целиком, а сколько ещё продастся сверх уже проданного. Для матчей, которые уже в продаже, отдельная модель учитывает, сколько осталось свободных мест».\n\n"
                 "Если спросят про абляции: «Без календаря RMSE 121,0, без соперника 124,1 — разница меньше погрешности валидации, поэтому мы не делаем вывод, что эти признаки не нужны».\n\n"
                 "Цифры абляций: без продаж на дату 299,3; без истории соперника 124,1; без КХЛ 122,5; без календаря 121,0 (при полной модели 124,7) ⚠. "
                 "Масштаб погрешности — бутстрэп со слайда 5: 95% ДИ разницы CatBoost − умный baseline шириной около 20 билетов (от −20,4 до −0,1), а изменения RMSE в абляциях без соперника, КХЛ и календаря — 0,5–3,7, то есть заметно меньше. Отдельного бутстрэпа для абляций нет.\n\n"
                 "Источники: src/baseline.py, src/model.py (докстринг), 03_model.ipynb → Абляции.",
    },
    {   # 5
        "layout": "text_top_image_table",
        "title": "Результат: CatBoost точнее baseline на 12%",
        "bullets": [
            "RMSE на проверке: CatBoost 124,7 ⚠, умный baseline 135,2, baseline из задания 141,0 (билетов на зону).",
            "По датам прогноза: 163,3 / 83,1 / 121,6 ⚠ против baseline 172,2 / 103,0 / 147,2. CatBoost лучше baseline из задания на всех трёх датах; умный baseline обходит его только на дате 25.11.2024 (79,1 против 83,1 ⚠).",
            "Против baseline из задания выигрыш надёжный: в среднем −16,4 билета на зону, 95% ДИ от −25,5 до −7,2; CatBoost лучше в 1 999 из 2 000 повторов бутстрэпа ⚠.",
            "Против умного baseline выигрыш есть, но на грани статистической значимости на 43 матчах: −10,5 билета, 95% ДИ от −20,4 до −0,1; CatBoost лучше в 1 953 из 2 000 повторов (≈98%) ⚠.",
            "На матч целиком CatBoost на уровне простого среднего (RMSE 447,5 против 450,9 билета на матч ⚠); выигрыш — в распределении по зонам (WAPE по строкам 13,2% против 16,0% ⚠).",
        ],
        "image": "01_rmse_models.png",
        "table": [
            ["CatBoost минус…", "Разница RMSE", "95% ДИ", "CatBoost лучше"],
            ["baseline из задания", "−16,4", "−25,5 … −7,2", "1 999 из 2 000"],
            ["умный baseline", "−10,5", "−20,4 … −0,1", "1 953 из 2 000"],
        ],
        "table_pt": 16,
        "notes": "«Мы сравнили с двумя ориентирами, а не с одним: обойти простое среднее легко, а умный baseline уже учитывает текущий сезон и проданное. Чтобы не выдавать удачу за качество, мы 2 000 раз пересобрали проверку из тех же 43 матчей. Простое среднее CatBoost обходит уверенно. Умный baseline — тоже почти всегда, но верхняя граница интервала почти касается нуля: на 43 матчах это на грани значимости, и мы так и говорим».\n\n"
                 "Как считали: бутстрэп по матчам, а не по строкам (8 зон матча всегда вместе), 2 000 повторов, random_state = 42, по прогнозам валидации из reports/cache/; модели не переобучались. 95% ДИ — перцентили 2,5 и 97,5.\n\n"
                 "Бизнес-метрики на матч целиком (сумма 8 зон, 43 матча проверки; baseline / умный baseline / CatBoost ⚠):\n"
                 "RMSE на матч: 450,9 / 519,4 / 447,5 билета\n"
                 "MAE на матч: 374,3 / 457,6 / 377,5 билета\n"
                 "Ошибка в % вместимости матча: 5,3 / 6,5 / 5,3\n"
                 "WAPE по матчам: 7,7% / 9,4% / 7,8%\n"
                 "WAPE по строкам матч × зона: 16,0% / 14,7% / 13,2%\n"
                 "Средний bias на матч (прогноз − факт): −177,6 / −330,4 / −217,3\n"
                 "«На матч целиком CatBoost не точнее простого среднего: общий спрос на матч задаёт уровень сезона. Выигрыш модели — в том, как она раскладывает спрос по зонам».\n\n"
                 "Источники: solution.ipynb → Валидация; 04_report.ipynb → 2.1, → Часть 3. Значимость выигрыша, → 4.5 (бизнес-метрики).",
    },
    {   # 6
        "layout": "text_top_image",
        "title": "Точность по зонам",
        "bullets": [
            "CatBoost лучше baseline в 6 зонах из 8 ⚠. Сильнее всего в зонах 4 (119,6 против 164,3) и 6 (202,6 против 257,0) ⚠.",
            "В дорогих зонах 1–3 и в зоне 8 ошибка небольшая: 26–55 билетов ⚠.",
            "Хуже baseline в зонах 5 (117,5 против 109,4) и 7 (218,8 против 208,5) ⚠.",
            "Почти вся ошибка сосредоточена в массовых зонах 6 и 7: RMSE 203 и 219 ⚠.",
        ],
        "image": "02_rmse_by_zone.png",
        "notes": "«Честно показываем, где модель не выигрывает: в зонах 5 и 7. Дальше объясним, почему».\n\n"
                 "RMSE по зонам (baseline / умный baseline / CatBoost ⚠):\n"
                 "Зона 1: 41,6 / 36,8 / 26,2\n"
                 "Зона 2: 38,4 / 28,9 / 26,2\n"
                 "Зона 3: 63,3 / 59,1 / 53,1\n"
                 "Зона 4: 164,3 / 139,0 / 119,6\n"
                 "Зона 5: 109,4 / 119,5 / 117,5\n"
                 "Зона 6: 257,0 / 225,7 / 202,6\n"
                 "Зона 7: 208,5 / 228,4 / 218,8\n"
                 "Зона 8: 58,8 / 62,2 / 55,5\n\n"
                 "Источники: 03_model.ipynb → По зонам; 04_report.ipynb → 0 (RMSE по зонам), 2.2.",
    },
    {   # 7
        "layout": "text_left_image",
        "title": "Разбор ошибок: зоны",
        "bullets": [
            "Самые крупные промахи — в зонах 6 и 7. Спрос между ними «переключается»: на одних матчах зона 6 продаёт около 1 000 билетов, а зона 7 около 350, на других наоборот (например, M029–M032).",
            "До открытия продаж такое переключение по нашим данным почти не предсказать. Когда продажи идут, модель видит его по уже проданному, поэтому на матчах в продаже ошибка 105,8, а на ещё не открытых 132,0 ⚠.",
            "Модель в среднем немного занижает: на 27 билетов на зону ⚠. Сильнее всего в зонах 4 и 5 (−73 и −79) ⚠.",
        ],
        "image": "03_fact_vs_pred.png",
        "notes": "«Точки далеко от диагонали — почти все из зон 6 и 7. Это не случайный шум, а перераспределение зрителей между соседними секторами. Без цен и схемы зала мы его объяснить не можем».\n\n"
                 "На графике: выбросы над и под диагональю — светлые точки зон 6 и 7; в заголовке RMSE 124,7 и корреляция 0,85 ⚠.\n\n"
                 "Источники: 03_model.ipynb → Разбор ошибок, → Матчи в продаже; 04_report.ipynb → 2.6 (общее смещение −27,2).",
    },
    {   # 8
        "layout": "text_top_image",
        "title": "Разбор ошибок: дни недели и соперники",
        "bullets": [
            "По дням недели CatBoost лучше baseline каждый день ⚠. Сильнее всего ошибаемся по четвергам (RMSE 166 ⚠), меньше всего по пятницам (95 ⚠).",
            "Четверг выделяется не из-за дня недели. Ошибку дают 3 матча из 6 (M030, M034, M084), и 64% её приходится на зоны 6 и 7 ⚠ — то самое переключение спроса со слайда 7.",
            "Модель занижает во все 7 дней недели ⚠: прогноз осторожный, а не перекошенный в какой-то день.",
            "По соперникам выводов не делаем: с каждым соперником на проверке было 1–5 матчей, у 13 соперников из 20 меньше трёх. Отдельные высокие ошибки (Сибирь 237, Спартак 195 ⚠) — это 1–2 конкретных матча.",
        ],
        "image": "06_errors_weekday.png",
        "notes": "«Задание просит разбор по соперникам и дням недели. Мы его сделали, но честно говорим: 43 матча на проверке слишком мало, чтобы утверждать “с Сибирью модель ошибается сильнее”. Видно одно: большие промахи — это всё те же зоны 6 и 7».\n\n"
                 "Таблица по соперникам — в 04_report.ipynb → 2.6, показать по вопросу.\n\n"
                 "Источник: 04_report.ipynb → 2.6 (таблицы по дням недели и соперникам, разбор четверга).",
    },
    {   # 9
        "layout": "text_top_image",
        "title": "Прогноз на 17 матчей и новогодние матчи",
        "bullets": [
            "Прогноз на 17 матчей — 85 631 билет ⚠ (5,0 тыс. за матч ⚠, +5% к сыгранной части сезона, 4 797); уже продано 16 106 на 6 матчей.",
            "Интервал на матч откалиброван на валидации (накрыл 38 из 43 матчей ⚠); по 17 матчам 71–105 тыс. билетов при прогнозе 85,6 тыс. ⚠",
            "Новогодние матчи 26–30.12 продаются раньше, но не больше: прошлые Новые годы закончили на 66–72% мест, модель ставит 73–77% ⚠ — риск скорее в сторону переоценки.",
        ],
        "image": "07_test_intervals.png",
        "notes": "Новогодние матчи 26–30.12 (Металлург Мг, Лада, Барыс) — ранний спрос, как в прошлые Новые годы. Продано уже 51–59% мест, в 1,6–2,1 раза быстрее обычного матча. Но прошлые новогодние матчи к той же дате были проданы на 58–68% (4 из 5) и закончили на 66–72%: спрос приходит раньше, но не больше. "
                 "В этом году к той же дате продано меньше, чем в прошлые Новые годы. По новогодней кривой итог был бы 56–63%, модель ставит 73–77% ⚠ — она учитывает, что сезон 2025/26 в целом заполненнее (70% против 63–66%). Аналогия построена на 5 матчах, поэтому риск отклонения в обе стороны, скорее в сторону переоценки. Прогноз модели не правили, чтобы не подгонять под 5 матчей.\n\n"
                 "Прошлые Новые годы (продано к той же дате, за 11–15 дней до матча, как у M087–M089 на 15.12 | итог):\n"
                 "M023, 24.12.2023: 58–62% | 70%\n"
                 "M024, 03.01.2024: 64–68% | 72%\n"
                 "M025, 06.01.2024: 62–66% | 71%\n"
                 "M055, 28.12.2024: 39–48% | 70%\n"
                 "M056, 30.12.2024: 60–62% | 66%\n"
                 "2025/26, M087–M089: 51–59% | прогноз 73–77% ⚠\n\n"
                 "Ни один матч и ни одна зона не выходят за исторические пределы и не упираются в число свободных мест ⚠.\n\n"
                 "«На новогодние матчи уже продано больше половины мест, и наши флаги на это срабатывают: обычный матч к этому дню продан вдвое меньше. Но прошлые Новые годы показывают, что это ранний спрос, а не больший: к той же дате четыре из пяти были проданы даже сильнее, чем сейчас, и все и закончили на 66–72%. Модель ставит выше, потому что сезон в целом заполненнее. Мы её не правили вручную — пять матчей слишком мало, чтобы подгонять под них прогноз, — но если ошибёмся, то скорее вверх».\n\n"
                 "На графике: столбик — уже продано на 15.12, точка — прогноз CatBoost, отрезок — интервал 10–90% на матч; жирные даты — матчи в продаже. "
                 "Соперники по датам: 24.12 Нефтехимик; 26.12 Металлург Мг; 28.12 Лада; 30.12 Барыс; 13.01 Салават Юлаев; 15.01 ЦСКА; 26.01 Адмирал; 29.01 Трактор; 05.02 Автомобилист; 15.02 Трактор; 17.02 Ак Барс; 20.02 Сибирь; 22.02 Амур; 27.02 Металлург Мг; 09.03 Сибирь; 11.03 Торпедо; 13.03 Шанхайские Драконы. "
                 "Интервал на матч — сумма зональных интервалов (квантили 10/90% остатков CatBoost на валидации по группам зон {1,2,8} / {3,4,5} / {6,7} и статусу продаж). "
                 "Честная проверка leave-one-fold-out: накрыл 38 из 43 матчей (88%) при номинале 80%. Ширина — 35% прогноза у матчей в продаже, 42% у остальных. "
                 "Прямой интервал по остатку суммы матча вдвое уже, но накрыл только 26 из 43 (60%): не учитывает межсезонный сдвиг. "
                 "Зональные интервалы — ориентир: покрытие 69% при номинале 80%. «71–105 тыс.» — сумма нижних и верхних границ по 17 матчам. "
                 "График без интервалов, с линией среднего сезона — reports/figures/04_test_forecast.png, показать по вопросу.\n\n"
                 "Источники: solution.ipynb → Сохранение (сумма); 04_report.ipynb → 1.1 (среднее сезона), 1.2 (темп и проданное), 1.3, 1.4 (флаги и пояснение), 1.5 (прошлые Новые годы), 2.4, 4.3–4.4 (интервалы).",
    },
    {   # 10
        "layout": "cards",
        "title": "Выводы и рекомендации клубу",
        "bullets": [
            "**Что сильнее всего определяет итог — темп уже начавшихся продаж.** Без него ошибка растёт с 125 до 299 ⚠. Рекомендация: пересчитывать прогноз каждую неделю по свежим продажам, особенно после открытия продаж на матч.",
            "**Новогодние матчи 26–30.12 — ранний спрос, а не аншлаг.** Продажи идут в 1,6–2,1 раза быстрее обычного, но так было и в прошлые Новые годы, а закончили они на 66–72% мест. Ориентир — прошлые Новые годы, а не темп первых недель; стоит следить за остатком мест и обновлять прогноз по свежим продажам.",
            "**Признаки темпа продаж проверили:** 3 варианта дают RMSE 122,6–124,6 против 124,7 у финальной модели ⚠, это в пределах шума (±3). Прогноз на новогодние матчи не меняется. Модель оставили без изменений.",
            "**Зона 6 — резерв выручки.** Заполняемость 39% при таком же числе свободных мест, как в зоне 7 (63%), а спрос между зонами 6 и 7 переключается от матча к матчу. Стоит проверить цену и позиционирование зоны 6 (цен в наших данных нет, это гипотеза для клуба).",
            "**Все модели в среднем занижают, но знак смещения меняется от сезона к сезону** (у CatBoost +76, −364 и −352 билета на матч по трём датам проверки ⚠), поэтому поправку не вводим.",
        ],
        "notes": "«Календарь и соперник, по нашим данным, влияют на спрос слабее, чем принято думать: их вклад в пределах погрешности проверки. Самая полезная информация — как продаются билеты прямо сейчас. Поэтому главное, что мы предлагаем клубу, — не разовый прогноз, а прогноз, который обновляется вместе с продажами».\n\n"
                 "Источник по признакам темпа: 03_model → Вывод: новогодние матчи и признак темпа. Смещение по датам проверки: 04_report → 4.5.",
    },
]

# ---------------------------------------------------------------- измерение текста
WARNINGS: list[str] = []
MOVED_TO_NOTES: list[str] = []
_FONTS: dict[bool, ImageFont.FreeTypeFont | None] = {}


def _font(bold: bool):
    if bold not in _FONTS:
        name = "calibrib.ttf" if bold else "calibri.ttf"
        cands = [Path("C:/Windows/Fonts") / name, Path("/usr/share/fonts/truetype/crosextra") / name]
        _FONTS[bold] = next((ImageFont.truetype(str(p), 100) for p in cands if p.exists()), None)
    return _FONTS[bold]


def clean(s: str) -> str:
    """Убирает пометки ⚠ вместе с пробелом перед ними; пробел между разрядами числа — неразрывный."""
    s = re.sub(r"\s*⚠", "", s)
    return re.sub(r"(?<=\d) (?=\d{3}\b)", "\u00a0", s)


def runs(s: str) -> list[tuple[str, bool]]:
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]."""
    parts = re.split(r"\*\*(.+?)\*\*", clean(s))
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


def text_width(s: str, pt: float, bold: bool = False) -> float:
    """Ширина строки в дюймах (по метрикам Calibri совпадает с переносами PowerPoint)."""
    f = _font(bold)
    if f is None:                         # нет Calibri — грубая оценка с запасом
        return len(s) * pt * 0.55 / 72
    return f.getlength(s) / 100 * pt / 72


def n_lines(s: str, width_in: float, pt: float) -> int:
    """Число строк при жадном переносе по словам; жирные фрагменты учитываются."""
    # делим только по обычному пробелу: неразрывный (\u00a0) держит число целиком, как в PowerPoint
    words = [(w, b) for chunk, b in runs(s) for w in re.findall(r"[^ \n]+ *", chunk)]
    lines, cur = 1, 0.0
    for w, b in words:
        ww = text_width(w, pt, b)
        if cur and cur + text_width(w.rstrip(), pt, b) > width_in:
            lines, cur = lines + 1, ww
        else:
            cur += ww
    return lines


LINE_K = 1.2                              # высота строки = 1.2 × кегль
SPACE_AFTER_PT = 8
BULLET_INDENT = 0.3


def bullets_height(items, width_in, pt=BODY_PT, bullet=True) -> float:
    w = width_in - (BULLET_INDENT if bullet else 0)
    h = sum(n_lines(s, w, pt) * pt * LINE_K / 72 for s in items)
    return h + SPACE_AFTER_PT / 72 * (len(items) - 1)


# ---------------------------------------------------------------- примитивы
def _rgb_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def add_runs(par, s: str, pt: float, color=INK, bold_all=False):
    for chunk, b in runs(s):
        r = par.add_run()
        r.text = chunk
        r.font.name = FONT
        r.font.size = Pt(pt)
        r.font.bold = b or bold_all
        r.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tb


def set_bullet(par, indent=BULLET_INDENT, color=ACCENT):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def add_bullets(slide, items, x, y, w, h, where: str, pt=BODY_PT):
    need = bullets_height(items, w, pt)
    if need > h + 0.02:
        WARNINGS.append(f"{where}: тезисы не помещаются — нужно {need:.2f}″, есть {h:.2f}″")
    tf = textbox(slide, x, y, w, h).text_frame
    for i, s in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_bullet(par)
        par.space_after = Pt(SPACE_AFTER_PT)
        par.line_spacing = 1.0
        add_runs(par, s, pt)
    return need


def add_title(slide, text):
    tf = textbox(slide, ML, TITLE_Y, CW, TITLE_H, anchor=MSO_ANCHOR.MIDDLE).text_frame
    add_runs(tf.paragraphs[0], text, TITLE_PT, color=ACCENT, bold_all=True)


def add_page_number(slide, n):
    tf = textbox(slide, W - MR - 1.0, 7.1, 1.0, 0.3).text_frame
    par = tf.paragraphs[0]
    par.alignment = PP_ALIGN.RIGHT
    add_runs(par, str(n), 12, color=MUTED)


def add_image_fit(slide, name, x, y, w, h, where):
    """Вписывает картинку в рамку с сохранением пропорций, по центру рамки."""
    path = FIG / name
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    if h < 2.0:
        WARNINGS.append(f"{where}: на картинку осталось мало высоты ({h:.2f}″)")
    slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             Inches(pw), Inches(ph))
    return pw, ph


CELL_MX, CELL_MY = 0.06, 0.04            # поля ячейки таблицы


def _cell(rows, i, j):
    """Текст ячейки для измерения: строка заголовка целиком жирная."""
    return f"**{rows[i][j]}**" if i == 0 else rows[i][j]


def row_heights(rows, col_w, pt):
    return [max(n_lines(_cell(rows, i, j), cw - 2 * CELL_MX, pt) for j, cw in enumerate(col_w))
            * pt * LINE_K / 72 + 2 * CELL_MY + 0.04 for i in range(len(rows))]


def table_height(rows, col_w, pt):
    return sum(row_heights(rows, col_w, pt))


def _widest(s, pt):
    return max(text_width(chunk, pt, b) for chunk, b in runs(s))


def _widest_word(s, pt):
    return max(text_width(w, pt, b) for chunk, b in runs(s) for w in chunk.split())


def col_widths(rows, total_w, pt):
    """Ширины по содержимому: обычные строки — в одну строку; заголовок и жирные строки
    переносятся только по пробелам (диапазон «73–77%» не рвётся)."""
    need = []
    for j in range(len(rows[0])):
        cells = [_cell(rows, i, j) for i in range(len(rows))]
        plain = [c for c in cells[1:] if "**" not in c]
        one_line = max((_widest(c, pt) for c in plain), default=0.0)
        words = max(_widest_word(c, pt) for c in cells)
        need.append(max(one_line, words) + 2 * CELL_MX + 0.05)
    k = total_w / sum(need)
    return [n * k for n in need]


def add_table(slide, rows, x, y, w, pt):
    cw = col_widths(rows, w, pt)
    rh = row_heights(rows, cw, pt)
    h = sum(rh)
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    for j, wj in enumerate(cw):
        tbl.columns[j].width = Inches(wj)
    for i, hi in enumerate(rh):
        tbl.rows[i].height = Inches(hi)
    for i, r in enumerate(rows):
        for j, txt in enumerate(r):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TINT if i == 0 else WHITE
            cell.margin_left = cell.margin_right = Inches(CELL_MX)
            cell.margin_top = cell.margin_bottom = Inches(CELL_MY)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            add_runs(par, txt, pt, color=INK, bold_all=(i == 0))
    return h


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = clean(text)


# ---------------------------------------------------------------- раскладки
def lay_text_top_image(slide, s, where):
    th = add_bullets(slide, s["bullets"], ML, TOP, CW, BOTTOM - TOP, where)
    y = TOP + th + GAP
    add_image_fit(slide, s["image"], ML, y, CW, BOTTOM - y, where)


def lay_text_top_image_table(slide, s, where):
    th = add_bullets(slide, s["bullets"], ML, TOP, CW, BOTTOM - TOP, where)
    y = TOP + th + GAP
    h = BOTTOM - y
    img_w = CW * 0.5
    pw, _ = add_image_fit(slide, s["image"], ML, y, img_w, h, where)
    # картинка прижата к левому краю, таблица — в оставшейся ширине справа
    slide.shapes[-1].left = Inches(ML)
    tx = ML + pw + GAP + 0.1
    tw = W - MR - tx
    rows, pt = s["table"], s.get("table_pt", 16)
    th_tab = table_height(rows, col_widths(rows, tw, pt), pt)
    add_table(slide, rows, tx, y + max(0.0, (h - th_tab) / 2), tw, pt)
    if th_tab > h:
        WARNINGS.append(f"{where}: таблица выше доступной области ({th_tab:.2f}″ > {h:.2f}″)")


def lay_text_left_image(slide, s, where):
    tw = CW * 0.52
    add_bullets(slide, s["bullets"], ML, TOP, tw, BOTTOM - TOP, where)
    ix = ML + tw + GAP + 0.2
    add_image_fit(slide, s["image"], ix, TOP, W - MR - ix, BOTTOM - TOP, where)


def lay_text_left_image_table(slide, s, where):
    tw = CW * 0.62
    add_bullets(slide, s["bullets"], ML, TOP, tw, BOTTOM - TOP, where)
    rx = ML + tw + GAP
    rw = W - MR - rx
    rows, pt = s["table"], max(TABLE_PT_MIN, s.get("table_pt", 14))
    th_tab = table_height(rows, col_widths(rows, rw, pt), pt)
    iw, ih = Image.open(FIG / s["image"]).size
    # картинку можно уменьшить (до 2″ по высоте), чтобы таблица встала под ней
    img_h = min(rw * ih / iw, BOTTOM - TOP - GAP - th_tab)
    if img_h >= 2.0:
        add_image_fit(slide, s["image"], rx, TOP, rw, img_h, where)
        add_table(slide, rows, rx, TOP + img_h + GAP, rw, pt)
    else:                                 # не влезает — таблица уходит в заметки
        add_image_fit(slide, s["image"], rx, TOP, rw, min(rw * ih / iw, BOTTOM - TOP), where)
        lines = [" | ".join(clean(c).replace("**", "") for c in r) for r in rows]
        s["notes"] += "\n\nТаблица прошлых Новых годов:\n" + "\n".join(lines)
        MOVED_TO_NOTES.append(f"{where}: таблица прошлых Новых годов (не поместилась при {pt} pt)")
    if s.get("table_note"):
        s["notes"] += "\n\n" + s["table_note"]


def lay_text_callouts(slide, s, where):
    tw = CW * 0.6
    add_bullets(slide, s["bullets"], ML, TOP, tw, BOTTOM - TOP, where)
    cx = ML + tw + GAP + 0.3
    cwid = W - MR - cx
    n = len(s["callouts"])
    box_h = 1.9
    total = n * box_h + (n - 1) * GAP
    y = TOP + (BOTTOM - TOP - total) / 2
    for big, label in s["callouts"]:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(y),
                                     Inches(cwid), Inches(box_h))
        _rgb_fill(box, TINT)
        box.line.fill.background()
        box.adjustments[0] = 0.08
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        add_runs(p1, big, 48, color=ACCENT, bold_all=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        add_runs(p2, label, 18, color=MUTED)
        y += box_h + GAP


def lay_text_top_diagram(slide, s, where):
    th = add_bullets(slide, s["bullets"], ML, TOP, CW, BOTTOM - TOP, where)
    y = TOP + th + GAP
    h = min(BOTTOM - y, 1.6)
    if h < 1.1:
        WARNINGS.append(f"{where}: на схему осталось мало высоты ({h:.2f}″)")
    y = y + (BOTTOM - y - h) / 2
    n = len(s["diagram"])
    arrow_w = 0.55
    bw = (CW - (n - 1) * (arrow_w + 0.2)) / n
    x = ML
    for i, (head, sub) in enumerate(s["diagram"]):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(bw), Inches(h))
        _rgb_fill(box, TINT)
        box.line.color.rgb = LINE
        box.adjustments[0] = 0.1
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        add_runs(p1, head, 20, color=ACCENT, bold_all=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        add_runs(p2, sub, 18, color=INK)
        need = (n_lines(head, bw - 0.3, 20) * 20 + n_lines(sub, bw - 0.3, 18) * 18) * LINE_K / 72
        if need > h - 0.1:
            WARNINGS.append(f"{where}: текст блока «{head}» не помещается в схему")
        x += bw
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 0.1), Inches(y + h / 2 - 0.22),
                                        Inches(arrow_w), Inches(0.44))
            _rgb_fill(ar, ACCENT)
            ar.line.fill.background()
            x += arrow_w + 0.2


def lay_cards(slide, s, where):
    pad, vpad = 0.2, 0.1                  # отступы текста в карточке: по бокам и сверху/снизу
    gap = 0.12
    inner = CW - 2 * pad
    hs = [bullets_height([b], inner, bullet=False) + 2 * vpad for b in s["bullets"]]
    avail = BOTTOM - TOP - gap * (len(hs) - 1)
    if sum(hs) > avail + 0.02:
        # оценка с запасом: карточки ужимаем до доступной высоты, а реальное переполнение проверяем глазами
        WARNINGS.append(f"{where}: по оценке карточкам нужно {sum(hs):.2f}″ при {avail:.2f}″ — "
                        "ужаты, проверьте рендер")
        hs = [h * avail / sum(hs) for h in hs]
    extra = max(0.0, (avail - sum(hs)) / len(hs))
    y = TOP
    for b, hh in zip(s["bullets"], hs):
        hh += extra
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ML), Inches(y),
                                     Inches(CW), Inches(hh))
        _rgb_fill(box, TINT)
        box.line.fill.background()
        box.adjustments[0] = 0.12 * min(1.0, 0.9 / hh)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(pad)
        tf.margin_top = tf.margin_bottom = Inches(vpad)
        par = tf.paragraphs[0]
        par.alignment = PP_ALIGN.LEFT
        par.line_spacing = 1.0
        add_runs(par, b, BODY_PT)
        y += hh + gap


LAYOUTS = {
    "text_top_image": lay_text_top_image,
    "text_top_image_table": lay_text_top_image_table,
    "text_left_image": lay_text_left_image,
    "text_left_image_table": lay_text_left_image_table,
    "text_callouts": lay_text_callouts,
    "text_top_diagram": lay_text_top_diagram,
    "cards": lay_cards,
}


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = TITLE_SLIDE
    tf = textbox(slide, 1.0, 2.2, W - 2.0, 1.6, anchor=MSO_ANCHOR.BOTTOM).text_frame
    add_runs(tf.paragraphs[0], t["title"], 44, color=ACCENT, bold_all=True)
    tf = textbox(slide, 1.0, 4.0, W - 2.0, 0.6).text_frame
    add_runs(tf.paragraphs[0], t["subtitle"], 26, color=INK)
    tf = textbox(slide, 1.0, 4.7, W - 2.0, 0.6).text_frame
    add_runs(tf.paragraphs[0], t["team"], 20, color=MUTED)
    return slide


# ---------------------------------------------------------------- сборка и PDF
def build() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    title_slide(prs)
    for i, s in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])     # пустой макет
        where = f"Слайд {i}"
        add_title(slide, s["title"])
        LAYOUTS[s["layout"]](slide, s, where)
        add_page_number(slide, i)
        set_notes(slide, s["notes"])
    prs.save(OUT_PPTX)
    return OUT_PPTX


def to_pdf(pptx: Path) -> Path | None:
    pdf = pptx.with_suffix(".pdf")
    if sys.platform == "win32":           # PowerPoint через COM
        ps = (
            "$ErrorActionPreference='Stop';"
            "$app = New-Object -ComObject PowerPoint.Application;"
            f"$p = $app.Presentations.Open('{pptx}', $true, $false, $false);"
            f"$p.SaveAs('{pdf}', 32); $p.Close(); $app.Quit()"
        )
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
        if r.returncode == 0 and pdf.exists():
            return pdf
        print("PowerPoint не сработал:", r.stderr.strip()[:300])
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                       check=False)
        if pdf.exists():
            return pdf
    return None


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    out = build()
    print(f"Сохранено: {out.relative_to(ROOT.parent)} (титульный + {len(SLIDES)} слайдов)")
    for m in MOVED_TO_NOTES:
        print("В заметки:", m)
    for w in WARNINGS:
        print("ВНИМАНИЕ:", w)
    if "--pdf" in sys.argv:
        pdf = to_pdf(out)
        print(f"PDF: {pdf.relative_to(ROOT.parent)}" if pdf else
              "PDF не создан: нет PowerPoint или LibreOffice — экспортируйте вручную.")
